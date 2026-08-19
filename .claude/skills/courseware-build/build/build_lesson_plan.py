#!/usr/bin/env python3
"""Generate the CompTIA Network+ (N10-009) Lesson Plan (LP) DOCX in the Tertiary house format.

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded 5-day schedule tables (9:30am-6:30pm, 8 training hours/day, 1h
lunch, tea within, final assessment Day 5 4:00pm). Topics/labs come from
course_data + the domain data files, and the Slides column is read from
slide_index.json (written by build_slides.py) so the LP always matches the deck.
"""
import os, sys, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

BRAND=RGBColor(0x1F,0x6F,0xEB); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"; ASSESS_FILL="E8F7EE"

# ---------------- slide index (written by build_slides.py) ----------------
IDX_PATH=os.path.join(HERE,"slide_index.json")
IDX=json.load(open(IDX_PATH)) if os.path.exists(IDX_PATH) else {"labs":{},"domains":{},"admin":{}}
LAB_SLIDE={int(k):v for k,v in IDX.get("labs",{}).items()}
DOM_SLIDE={int(k):v for k,v in IDX.get("domains",{}).items()}
LAST_SLIDE=max(list(LAB_SLIDE.values())+list(DOM_SLIDE.values())+[0])

def lab_range(lo,hi):
    """Slide range covering labs lo..hi (from first lab overview to the slide
    before the next lab / domain section)."""
    if lo not in LAB_SLIDE: return "—"
    start=LAB_SLIDE[lo]
    nxt=[v for v in list(LAB_SLIDE.values())+list(DOM_SLIDE.values()) if v>LAB_SLIDE.get(hi,start)]
    end=(min(nxt)-1) if nxt else LAST_SLIDE
    return f"Slides {start}–{end}"

def dom_range(n):
    if n not in DOM_SLIDE: return "—"
    start=DOM_SLIDE[n]
    first_lab=min([a["num"] for a in ACT if a["topic"]==n])
    end=LAB_SLIDE.get(first_lab,start+1)-1
    return f"Slides {start}–{end}"

def lab_titles(nums):
    return "; ".join(f"Lab {a['num']}: {a['title']}" for a in ACT if a['num'] in nums)

# ------------------------------------------------ schedule (single source of truth for timing)
# rows: (start, end, minutes, kind, activity_text, slides)
SCHEDULE = {
 1: ("Networking Concepts", [
    ("9:30","10:00",30,"admin","Welcome, course introduction, learning outcomes, ground rules and mandatory digital attendance (AM)","—"),
    ("10:00","11:00",60,"topic","Domain 1.0 — Networking Concepts: OSI model layers, encapsulation and PDUs; networking appliances and their functions",dom_range(1)),
    ("11:00","11:15",15,"break","Tea break","—"),
    ("11:15","13:00",105,"lab","Concepts + hands-on: "+lab_titles([1,2]),lab_range(1,2)),
    ("13:00","14:00",60,"lunch","Lunch break","—"),
    ("14:00","15:45",105,"lab","Concepts + hands-on: "+lab_titles([3,4]),lab_range(3,4)),
    ("15:45","16:00",15,"break","Tea break","—"),
    ("16:00","18:15",135,"lab","Concepts + hands-on: "+lab_titles([5,6,7]),lab_range(5,7)),
    ("18:15","18:30",15,"recap","Day 1 recap, Q&A and PM digital attendance","—"),
 ]),
 2: ("Network Implementation", [
    ("9:30","9:45",15,"recap","Day 1 recap and mandatory digital attendance (AM)","—"),
    ("9:45","11:00",75,"topic","Domain 2.0 — Network Implementation: routing technologies (static vs dynamic, BGP/EIGRP/OSPF, route selection, NAT/PAT, FHRP)",dom_range(2)),
    ("11:00","11:15",15,"break","Tea break","—"),
    ("11:15","13:00",105,"lab","Concepts + hands-on: "+lab_titles([8,9]),lab_range(8,9)),
    ("13:00","14:00",60,"lunch","Lunch break","—"),
    ("14:00","15:45",105,"lab","Concepts + hands-on: "+lab_titles([10,11]),lab_range(10,11)),
    ("15:45","16:00",15,"break","Tea break","—"),
    ("16:00","18:15",135,"lab","Concepts + hands-on: "+lab_titles([12,13,14]),lab_range(12,14)),
    ("18:15","18:30",15,"recap","Day 2 recap, Q&A and PM digital attendance","—"),
 ]),
 3: ("Network Operations", [
    ("9:30","9:45",15,"recap","Day 2 recap and mandatory digital attendance (AM)","—"),
    ("9:45","11:00",75,"topic","Domain 3.0 — Network Operations: documentation, IPAM, life-cycle, change and configuration management",dom_range(3)),
    ("11:00","11:15",15,"break","Tea break","—"),
    ("11:15","13:00",105,"lab","Concepts + hands-on: "+lab_titles([15,16]),lab_range(15,16)),
    ("13:00","14:00",60,"lunch","Lunch break","—"),
    ("14:00","15:45",105,"lab","Concepts + hands-on: "+lab_titles([17,18]),lab_range(17,18)),
    ("15:45","16:00",15,"break","Tea break","—"),
    ("16:00","18:15",135,"lab","Concepts + hands-on: "+lab_titles([19,20])+". Domain 3 recap and consolidation",lab_range(19,20)),
    ("18:15","18:30",15,"recap","Day 3 recap, Q&A and PM digital attendance","—"),
 ]),
 4: ("Network Security", [
    ("9:30","9:45",15,"recap","Day 3 recap and mandatory digital attendance (AM)","—"),
    ("9:45","11:00",75,"topic","Domain 4.0 — Network Security: encryption, PKI, IAM/MFA/SSO, AAA (RADIUS/TACACS+), security terminology and the CIA triad",dom_range(4)),
    ("11:00","11:15",15,"break","Tea break","—"),
    ("11:15","13:00",105,"lab","Concepts + hands-on: "+lab_titles([21,22]),lab_range(21,22)),
    ("13:00","14:00",60,"lunch","Lunch break","—"),
    ("14:00","15:45",105,"lab","Concepts + hands-on: "+lab_titles([23,24]),lab_range(23,24)),
    ("15:45","16:00",15,"break","Tea break","—"),
    ("16:00","18:15",135,"lab","Concepts + hands-on: "+lab_titles([25,26,27]),lab_range(25,27)),
    ("18:15","18:30",15,"recap","Day 4 recap, Q&A and PM digital attendance","—"),
 ]),
 5: ("Network Troubleshooting & Assessment", [
    ("9:30","9:45",15,"recap","Day 4 recap and mandatory digital attendance (AM)","—"),
    ("9:45","11:00",75,"topic","Domain 5.0 — Network Troubleshooting: the CompTIA seven-step methodology; cabling and physical interface issues",dom_range(5)),
    ("11:00","11:15",15,"break","Tea break","—"),
    ("11:15","13:00",105,"lab","Concepts + hands-on: "+lab_titles([28,29,30,31]),lab_range(28,31)),
    ("13:00","14:00",60,"lunch","Lunch break","—"),
    ("14:00","15:30",90,"lab","Concepts + hands-on: "+lab_titles([32,33,34,35,36])+". Course recap and exam preparation",lab_range(32,36)),
    ("15:30","15:45",15,"break","Tea break","—"),
    ("15:45","16:00",15,"assess","Briefing for Assessment","—"),
    ("16:00","17:00",60,"assess","Written Assessment (WA) — Short-Answer Questions (SAQ), 1 hour, open book","—"),
    ("17:00","18:30",90,"assess","Practical Performance (PP) — hands-on networking tasks, 90 minutes, open book. PM digital attendance","—"),
 ]),
}

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=os.path.join(ASSETS,"comptia-network-badge.png"),
                      course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("1.0","2 July 2026","Initial release — Network+ (N10-009) 5-day lesson plan aligned to the 36 hands-on labs.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Regenerated against the v2 visual slide deck (five N10-009 domains, per-lab folders, TIS web "
  "ed-tools: IP Calculator, PCAP Analyzer, Cybersecurity Simulator, Regex Generator). Slide "
  "references updated to the new deck.",C.TRAINER),
])
prodoc.add_toc(doc)

def H(text,level=1):
    return doc.add_heading(text,level=level)

H("Course Information",1)
info=[("Course Title",C.TITLE),("WSQ Course Reference",C.COURSE_CODE),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("Duration","5 days · 8 training hours per day (40 hours)"),
      ("Daily Timing","9:30 am – 6:30 pm (1-hour lunch; tea breaks within training time)"),
      ("Mode","Instructor-led, hands-on labs on the Killercoda Ubuntu Playground and free TIS web tools"),
      ("Trainer",C.TRAINER)]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("Assessment",1)
for a in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
          "Format: Open Book — course slides, Learner Guide and approved materials only.",
          "Final assessment is conducted on Day 5 from 4:00 pm (WA 4:00–5:00 pm, PP 5:00–6:30 pm).",
          C.ASSESSMENT["note"],
          "Courseware and the assessment are on the LMS (https://lms-tms.tertiaryinfotech.com/)."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None):
    cell.text=""; p=cell.paragraphs[0]
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","recap":"F3F5F8","lab":None}

H("Course Schedule",1)
doc.add_paragraph("The Slides column gives the matching page range in the course deck "
                  f"({C.SHORT_TITLE}-{C.VERSION}.pptx) so the trainer can run each session directly from the deck.")
for day,(theme,rows) in SCHEDULE.items():
    H(f"Day {day} — {theme}",2)
    tbl=doc.add_table(rows=0,cols=4); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tbl.add_row().cells
    for i,htext in enumerate(["Time","Duration","Topic / Activity","Slides"]):
        set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
    training=0
    for start,end,mins,kind,text,slides in rows:
        cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
        set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess")),fill=fill)
        set_cell(cells[1],f"{mins} min",fill=fill)
        set_cell(cells[2],text,bold=(kind in ("topic","assess")),fill=fill)
        set_cell(cells[3],slides,fill=fill)
        if kind!="lunch": training+=mins
    for row in tbl.rows:
        row.cells[0].width=Inches(1.0); row.cells[1].width=Inches(0.8)
        row.cells[2].width=Inches(4.0); row.cells[3].width=Inches(1.1)
    p=doc.add_paragraph(); r=p.add_run(f"Total training time: {training} minutes ({training//60} hours).")
    r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    assert training==480, f"Day {day} training minutes = {training}, expected 480"

H("Topic-by-Topic Breakdown",1)
doc.add_paragraph("Each domain and lab below shows its starting slide in the course deck.")
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    h=H(f"Domain {tp['code']} — {tp['title']}  ({tp['weighting']})   ·   {dom_range(tp['num']).replace('Slides','Slides')}",2)
    doc.add_paragraph(tp["subtitle"])
    for a in acts:
        sl=LAB_SLIDE.get(a["num"])
        p=doc.add_paragraph()
        r=p.add_run(f"Lab {a['num']} — {a['title']}   ·   {'Slide '+str(sl) if sl else '—'}")
        r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=BRAND
        p2=doc.add_paragraph(); p2.add_run(f"Objective: {a['objective']}").font.size=Pt(10)
        p3=doc.add_paragraph(); p3.add_run(f"Activity: {a['desc']}").font.size=Pt(10)

H("Lab Reference (aligned to exam domains)",1)
tt=doc.add_table(rows=0,cols=3); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Exam domain","Weighting","Labs"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Domain {tp['code']}: {tp['title']}",bold=True,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],fill=TOPIC_FILL)
    set_cell(cells[2],", ".join(f"Lab {a['num']}" for a in acts))

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT)
